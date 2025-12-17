from flask import Flask, request, jsonify, send_file
import os
import io
import requests
from bs4 import BeautifulSoup
from docx import Document
import google.generativeai as genai
import json
from flask_cors import CORS
import streamlit as st

app = Flask(__name__)
CORS(app)  # allow Streamlit (different port) to call this

# ---- Gemini setup ----
genai.configure(api_key=st.secrets["GOOGLE_API_KEY"])
model = genai.GenerativeModel("models/gemma-3n-e2b-it")  # or gemini-1.5-pro


# ---- Helper: extract text from webpage ----
def extract_page_text(url: str) -> str:
    resp = requests.get(url, headers={"User-Agent": "Mozilla/5.0"})
    resp.raise_for_status()
    html = resp.text

    soup = BeautifulSoup(html, "html.parser")

    # SIMPLE VERSION: take text from main tags
    texts = []

    # try common content containers first (you can tune this later)
    candidates = soup.select("article, .content, .post-content, .entry-content")
    if candidates:
        for c in candidates:
            texts.append(c.get_text(separator="\n", strip=True))
    else:
        # fallback: all p + headings + list items
        for tag in soup.find_all(["h1", "h2", "h3", "p", "li"]):
            txt = tag.get_text(separator=" ", strip=True)
            if txt:
                texts.append(txt)

    page_text = "\n".join(texts)
    # optional: cut very large pages
    max_chars = 30000
    return page_text[:max_chars]


# ---- Route 1: scrape ----
@app.post("/scrape")
def scrape():
    data = request.get_json()
    url = data.get("url")
    if not url:
        return jsonify({"error": "url is required"}), 400

    try:
        page_text = extract_page_text(url)
        return jsonify({"pageText": page_text})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
# Helper to parse LLM JSON output safely
# def parse_llm_json(raw_text: str):
#     import re, json

#     # Remove markdown fences
#     raw_text = re.sub(r"```json|```", "", raw_text).strip()

#     # Extract JSON array
#     match = re.search(r"\[\s*{.*}\s*\]", raw_text, re.DOTALL)
#     if not match:
#         raise ValueError("No JSON array found")

#     json_text = match.group()

#     # Normalize multiline strings
#     json_text = json_text.replace("\n", "\\n").replace("\t", "\\t")

#     return json.loads(json_text)

# ---- Route 2: answer questions with Gemini ----
@app.post("/answer-questions")
def answer_questions():
    data = request.get_json()
    page_text = data.get("pageText", "")
    questions = data.get("questions", [])

    if not page_text or not questions:
        return jsonify({"error": "pageText and questions are required"}), 400

    # Prepare prompt
    questions_block = "\n".join(
        [f"{i+1}. {q}" for i, q in enumerate(questions)]
    )

    prompt = f"""
You are an exam-notes assistant.

You are given the text of a single web page (PAGE_TEXT). 
You must answer each question ONLY using this text. 
If the answer is not clearly available, mark found = false.

Return a JSON array. Each element must be:
{{
  "question": "...",
  "answer": "...",
  "found": true or false,
  "source_snippet": "copy the exact lines or phrases you used from PAGE_TEXT, if any"
}}

Rules:
- Answer length: about 150–250 words if found do not need to add points just make it longer .
- Do NOT add knowledge from outside PAGE_TEXT.
- Every Question must be worth 15 marks.
- If not found, answer: "Not clearly available in the given page."

PAGE_TEXT:
\"\"\"{page_text}\"\"\"

QUESTIONS:
{questions_block}
"""

    try:
        response = model.generate_content(prompt)
        if not response.text:
            return jsonify({"error": "Empty response from Gemini"}), 500
        raw_text = response.text
        print("RAW GEMINI OUTPUT:\n", raw_text)
        import re, json
        cleaned = re.sub(r"```json|```", "", raw_text).strip()
        # The model should return JSON. Try to parse:
        qa_list = json.loads(cleaned)
        # qa_list = parse_llm_json(cleaned)

        return jsonify({"qa": qa_list})

        # return jsonify({"qa": qa_list})
    except Exception as e:
        print("ERROR:", str(e))
        # debug: you might want to log raw_text when parsing fails
        return jsonify({"error": str(e)}), 500


# ---- Route 3: export DOCX ----
@app.post("/export-docx")
def export_docx():
    data = request.get_json()
    qa_list = data.get("qa", [])
    title = data.get("title", "Study Notes")

    doc = Document()
    doc.add_heading(title, level=1)

    for item in qa_list:
        q = item.get("question", "")
        a = item.get("answer", "")
        found = item.get("found", True)

        doc.add_heading(q, level=2)
        if not found:
            doc.add_paragraph("(Not clearly available in the given page.)")
        doc.add_paragraph(a)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        as_attachment=True,
        download_name=f"{title.replace(' ', '_')}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
@app.route("/generate-questions", methods=["POST"])
def generate_questions():
    try:
        data = request.get_json(force=True)
        page_text = data.get("pageText", "")
        num_questions = data.get("num_questions", 5)
        marks = data.get("marks", 10)
        difficulty = data.get("difficulty", "exam")

        if not page_text:
            return jsonify({"error": "pageText is required"}), 400

        prompt = f"""
Return ONLY valid JSON.
Do NOT add markdown or ```.

Generate {num_questions} university exam-oriented questions
from the given PAGE_TEXT.

Guidelines:
- Questions must be suitable for {marks}-mark answers
- Difficulty: {difficulty}
- Use clear, direct exam-style wording
- Do NOT add answers
- do not add any special characters just return simple questiosn
- do not add ( ) or any brackets in the questions

Return format:
[
  "Question 1",
  "Question 2",
  "Question 3"
]

PAGE_TEXT:
\"\"\"{page_text}\"\"\"
"""

        response = model.generate_content(prompt)

        if not response.text:
            return jsonify({"error": "Empty response from Gemini"}), 500

        raw_text = response.text
        print("RAW GEMINI QUESTIONS OUTPUT:\n", raw_text)

        import re, json

        # Extract JSON array safely
        match = re.search(r"\[.*\]", raw_text, re.DOTALL)
        if not match:
            return jsonify({"error": "No JSON array found in Gemini output"}), 500

        questions = json.loads(match.group())

        return jsonify({"questions": questions})

    except Exception as e:
        print("ERROR generating questions:", str(e))
        return jsonify({
            "error": "Failed to generate questions",
            "details": str(e)
        }), 500



if __name__ == "__main__":
    app.run(debug=True, port=5000)
